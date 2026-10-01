import customtkinter as ctk
from tkinter import filedialog, messagebox
import threading
import queue
import os
import json
import time
import sys
import urllib.request
import shutil
import tempfile
import subprocess
from datetime import datetime

import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.options import Options
from webdriver_manager.chrome import ChromeDriverManager

VERSION = "1.0.0"
GITHUB_VERSION_URL = "https://raw.githubusercontent.com/Choi-Yu-Rim/CashReceiptAutoUpload/master/version.txt"
GITHUB_SCRIPT_URL  = "https://raw.githubusercontent.com/Choi-Yu-Rim/CashReceiptAutoUpload/master/main.py"

CONFIG_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.json")
ROWS_PER_FILE = 200
LOGIN_URL = "https://taxadmin.tosspayments.com/"
UPLOAD_URL = "https://taxadmin.tosspayments.com/receipt/receiptAdmitReq.jsp?pageNum=2&subNum=3"


def fetch_latest_version() -> str | None:
    """GitHub에서 최신 버전 문자열을 가져옴. 실패 시 None 반환."""
    try:
        with urllib.request.urlopen(GITHUB_VERSION_URL, timeout=5) as r:
            return r.read().decode().strip()
    except Exception:
        return None


def download_update(dest_path: str) -> bool:
    """최신 main.py를 dest_path에 저장. 성공 시 True 반환."""
    try:
        with urllib.request.urlopen(GITHUB_SCRIPT_URL, timeout=15) as r:
            content = r.read()
        with open(dest_path, "wb") as f:
            f.write(content)
        return True
    except Exception:
        return False


UPDATE_DONE_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".update_done.json")


def apply_update_and_restart(new_script: str, new_version: str):
    """현재 스크립트를 새 버전으로 교체하고 재실행."""
    current = os.path.abspath(__file__)
    shutil.copy2(new_script, current)
    # 재시작 후 팝업을 위해 업데이트 정보 저장
    with open(UPDATE_DONE_FILE, "w", encoding="utf-8") as f:
        json.dump({
            "version": new_version,
            "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }, f, ensure_ascii=False)
    subprocess.Popen([sys.executable] + sys.argv)
    sys.exit(0)

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")


def load_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {"user_id": "", "user_pw": ""}


def save_config(user_id, user_pw):
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump({"user_id": user_id, "user_pw": user_pw}, f, ensure_ascii=False)


def split_excel(file_path: str, output_dir: str):
    os.makedirs(output_dir, exist_ok=True)

    with open(file_path, "rb") as f:
        magic = f.read(8)

    if magic[:4] == b'PK\x03\x04':
        # .xlsx 포맷 (data_only=True: 수식 대신 마지막 계산된 값을 읽어 분할 파일에 수식 오류 방지)
        with open(file_path, "rb") as f:
            wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
    elif magic[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
        # .xls 포맷 (구형)
        import xlrd
        xls_wb = xlrd.open_workbook(file_path)
        xls_ws = xls_wb.sheet_by_index(0)
        all_rows = [tuple(xls_ws.row_values(i)) for i in range(xls_ws.nrows)]
    else:
        raise ValueError("파일을 읽을 수 없습니다. DRM 보안이 걸려 있거나 손상된 파일입니다.")

    if len(all_rows) < 2:
        raise ValueError("엑셀 파일에 데이터가 없습니다 (헤더 외 행이 필요합니다).")
    header = all_rows[0]
    data_rows = all_rows[1:]
    chunks = [data_rows[i:i + ROWS_PER_FILE] for i in range(0, len(data_rows), ROWS_PER_FILE)]
    output_files = []
    for idx, chunk in enumerate(chunks, 1):
        new_wb = openpyxl.Workbook()
        new_ws = new_wb.active
        new_ws.append(list(header))
        for row in chunk:
            new_ws.append(list(row))
        out_path = os.path.join(output_dir, f"현금영수증_분할_{idx}.xlsx")
        new_wb.save(out_path)
        output_files.append(out_path)
    return output_files


class App(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("현금영수증 자동 업로드")
        self.geometry("660x860")
        self.minsize(620, 500)
        self.resizable(True, True)

        self.config_data = load_config()
        self.msg_queue: queue.Queue = queue.Queue()
        self.running = False
        self.split_files: list = []

        self._build_ui()
        self._poll_queue()
        self.after(300, self._show_update_done)
        threading.Thread(target=self._check_update, daemon=True).start()

    def _build_ui(self):
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)

        sc = ctk.CTkScrollableFrame(self, fg_color="transparent")
        sc.grid(row=0, column=0, sticky="nsew")
        sc.grid_columnconfigure(0, weight=1)

        # 제목
        title_frame = ctk.CTkFrame(sc, fg_color="transparent")
        title_frame.grid(row=0, column=0, padx=24, pady=(24, 0), sticky="ew")
        ctk.CTkLabel(title_frame, text="현금영수증 자동 업로드",
                     font=ctk.CTkFont(size=22, weight="bold")).pack(anchor="w")
        ctk.CTkLabel(title_frame, text="토스페이먼츠 현금영수증 일괄 업로드 도우미",
                     font=ctk.CTkFont(size=12), text_color="gray").pack(anchor="w", pady=(2, 0))

        # ── Step 1. DRM 보안 해제 확인 ──
        f1 = ctk.CTkFrame(sc)
        f1.grid(row=1, column=0, padx=24, pady=(16, 0), sticky="ew")
        f1.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(f1, text="Step 1.  DRM 보안 해제 확인",
                     font=ctk.CTkFont(size=13, weight="bold")).grid(
            row=0, column=0, padx=16, pady=(14, 8), sticky="w")

        drm_box = ctk.CTkFrame(f1, fg_color="#fff8e1", corner_radius=8)
        drm_box.grid(row=1, column=0, padx=16, pady=(0, 14), sticky="ew")
        drm_box.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(drm_box,
                     text="⚠️  원본 파일에 DRM 보안이 걸려 있으면 파일 분할이 불가합니다.",
                     font=ctk.CTkFont(size=12), text_color="#e65100", anchor="w").grid(
            row=0, column=0, padx=14, pady=(10, 4), sticky="w")
        ctk.CTkLabel(drm_box,
                     text="아래 순서대로 보안을 해제한 뒤 진행해주세요.",
                     font=ctk.CTkFont(size=11), text_color="#555555", anchor="w").grid(
            row=1, column=0, padx=14, pady=(0, 4), sticky="w")
        ctk.CTkLabel(drm_box,
                     text="  1. DRM 프로그램이 설치된 회사 PC에서 원본 파일 열기\n"
                          "  2. 파일 → 다른 이름으로 저장 → Excel 통합 문서 (*.xlsx) 선택 후 저장\n"
                          "  3. 저장된 파일을 이 프로그램에서 사용",
                     font=ctk.CTkFont(size=11), text_color="#555555", anchor="w", justify="left").grid(
            row=2, column=0, padx=14, pady=(0, 12), sticky="w")

        # ── Step 2. 파일 선택 ──
        f2 = ctk.CTkFrame(sc)
        f2.grid(row=2, column=0, padx=24, pady=(12, 0), sticky="ew")
        f2.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(f2, text="Step 2.  파일 선택",
                     font=ctk.CTkFont(size=13, weight="bold")).grid(
            row=0, column=0, columnspan=2, padx=16, pady=(14, 10), sticky="w")

        ctk.CTkLabel(f2, text="원본 파일", font=ctk.CTkFont(size=11),
                     text_color="gray").grid(row=1, column=0, columnspan=2, padx=16, pady=(0, 4), sticky="w")
        src_row = ctk.CTkFrame(f2, fg_color="transparent")
        src_row.grid(row=2, column=0, columnspan=2, padx=16, pady=(0, 10), sticky="ew")
        src_row.grid_columnconfigure(0, weight=1)
        self.file_var = ctk.StringVar()
        ctk.CTkEntry(src_row, textvariable=self.file_var, state="readonly",
                     font=ctk.CTkFont(size=11), height=34).grid(row=0, column=0, padx=(0, 8), sticky="ew")
        ctk.CTkButton(src_row, text="찾기", width=70, height=34,
                      command=self._browse_src).grid(row=0, column=1)

        ctk.CTkLabel(f2, text="분할 파일 저장 폴더", font=ctk.CTkFont(size=11),
                     text_color="gray").grid(row=3, column=0, columnspan=2, padx=16, pady=(0, 4), sticky="w")
        out_row = ctk.CTkFrame(f2, fg_color="transparent")
        out_row.grid(row=4, column=0, columnspan=2, padx=16, pady=(0, 14), sticky="ew")
        out_row.grid_columnconfigure(0, weight=1)
        self.out_dir_var = ctk.StringVar()
        ctk.CTkEntry(out_row, textvariable=self.out_dir_var, state="readonly",
                     font=ctk.CTkFont(size=11), height=34).grid(row=0, column=0, padx=(0, 8), sticky="ew")
        ctk.CTkButton(out_row, text="변경", width=70, height=34,
                      command=self._browse_out).grid(row=0, column=1)

        # ── Step 3. 파일 분할 ──
        f3 = ctk.CTkFrame(sc)
        f3.grid(row=3, column=0, padx=24, pady=(12, 0), sticky="ew")
        f3.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(f3, text="Step 3.  파일 분할",
                     font=ctk.CTkFont(size=13, weight="bold")).grid(
            row=0, column=0, columnspan=2, padx=16, pady=(14, 10), sticky="w")

        split_bottom = ctk.CTkFrame(f3, fg_color="transparent")
        split_bottom.grid(row=1, column=0, columnspan=2, padx=16, pady=(0, 8), sticky="ew")
        self.split_btn = ctk.CTkButton(split_bottom, text="파일 분할", width=110, height=36,
                                       fg_color="#546e7a", hover_color="#37474f",
                                       command=self._start_split)
        self.split_btn.pack(side="left")
        self.split_result_var = ctk.StringVar(value="")
        ctk.CTkLabel(split_bottom, textvariable=self.split_result_var,
                     font=ctk.CTkFont(size=12)).pack(side="left", padx=(14, 0))

        # 분할 결과 정보 박스 (분할 후 표시)
        self.split_info_frame = ctk.CTkFrame(f3, fg_color="#e8f5e9", corner_radius=8)
        self.split_info_frame.grid(row=2, column=0, columnspan=2, padx=16, pady=(0, 14), sticky="ew")
        self.split_info_frame.grid_columnconfigure(0, weight=1)
        self.split_info_frame.grid_remove()

        info_top = ctk.CTkFrame(self.split_info_frame, fg_color="transparent")
        info_top.grid(row=0, column=0, padx=12, pady=(10, 10), sticky="ew")
        info_top.grid_columnconfigure(0, weight=1)

        self.info_folder_var = ctk.StringVar(value="")
        ctk.CTkLabel(info_top, textvariable=self.info_folder_var,
                     font=ctk.CTkFont(size=11), text_color="#1b5e20", anchor="w").grid(
            row=0, column=0, sticky="ew")
        self.open_folder_btn = ctk.CTkButton(
            info_top, text="폴더 열기", width=80, height=26,
            fg_color="#2e7d32", hover_color="#1b5e20",
            font=ctk.CTkFont(size=11), command=self._open_split_folder
        )
        self.open_folder_btn.grid(row=0, column=1, padx=(8, 0))

        # ── Step 4. 업로드 ──
        f4 = ctk.CTkFrame(sc)
        f4.grid(row=4, column=0, padx=24, pady=(12, 0), sticky="ew")
        f4.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(f4, text="Step 4.  업로드",
                     font=ctk.CTkFont(size=13, weight="bold")).grid(
            row=0, column=0, columnspan=2, padx=16, pady=(14, 10), sticky="w")

        # 로그인 정보
        ctk.CTkLabel(f4, text="아이디", font=ctk.CTkFont(size=12), width=60).grid(
            row=1, column=0, padx=(16, 8), pady=4, sticky="w")
        self.id_var = ctk.StringVar(value=self.config_data.get("user_id", ""))
        ctk.CTkEntry(f4, textvariable=self.id_var, height=36,
                     font=ctk.CTkFont(size=12)).grid(row=1, column=1, padx=(0, 16), pady=4, sticky="ew")

        ctk.CTkLabel(f4, text="비밀번호", font=ctk.CTkFont(size=12), width=60).grid(
            row=2, column=0, padx=(16, 8), pady=4, sticky="w")
        self.pw_var = ctk.StringVar(value=self.config_data.get("user_pw", ""))
        ctk.CTkEntry(f4, textvariable=self.pw_var, show="●", height=36,
                     font=ctk.CTkFont(size=12)).grid(row=2, column=1, padx=(0, 16), pady=4, sticky="ew")

        has_saved = bool(self.config_data.get("user_id") or self.config_data.get("user_pw"))
        saved_text = "✅  저장된 설정을 불러왔습니다." if has_saved else "⚠️  저장된 설정이 없습니다. 입력해주세요."
        ctk.CTkLabel(f4, text=saved_text, font=ctk.CTkFont(size=11),
                     text_color="#2e7d32" if has_saved else "#b45309").grid(
            row=3, column=0, columnspan=2, padx=16, pady=(4, 12), sticky="w")

        # 구분선
        ctk.CTkFrame(f4, height=1, fg_color="#e0e0e0").grid(
            row=4, column=0, columnspan=2, padx=16, pady=(0, 12), sticky="ew")

        # 업로드 폴더 선택
        ctk.CTkLabel(f4, text="업로드할 파일 폴더", font=ctk.CTkFont(size=11),
                     text_color="gray").grid(row=5, column=0, columnspan=2, padx=16, pady=(0, 4), sticky="w")
        upload_dir_row = ctk.CTkFrame(f4, fg_color="transparent")
        upload_dir_row.grid(row=6, column=0, columnspan=2, padx=16, pady=(0, 10), sticky="ew")
        upload_dir_row.grid_columnconfigure(0, weight=1)

        self.upload_dir_var = ctk.StringVar(value="")
        ctk.CTkEntry(upload_dir_row, textvariable=self.upload_dir_var, state="readonly",
                     font=ctk.CTkFont(size=11), height=34).grid(row=0, column=0, padx=(0, 8), sticky="ew")
        ctk.CTkButton(upload_dir_row, text="선택", width=60, height=34,
                      command=self._load_existing_folder).grid(row=0, column=1, padx=(0, 8))
        self.step3_open_btn = ctk.CTkButton(
            upload_dir_row, text="폴더 열기", width=80, height=34,
            fg_color="#2e7d32", hover_color="#1b5e20",
            font=ctk.CTkFont(size=11), state="disabled",
            command=self._open_split_folder
        )
        self.step3_open_btn.grid(row=0, column=2)

        # 범위 설정
        self.range_all_var = ctk.BooleanVar(value=True)
        ctk.CTkCheckBox(f4, text="전체 파일 업로드",
                        variable=self.range_all_var, command=self._toggle_range,
                        font=ctk.CTkFont(size=12)).grid(
            row=7, column=0, columnspan=2, padx=16, pady=(0, 8), sticky="w")

        range_row = ctk.CTkFrame(f4, fg_color="transparent")
        range_row.grid(row=8, column=0, columnspan=2, padx=16, pady=(0, 14), sticky="w")

        ctk.CTkLabel(range_row, text="시작 번호", font=ctk.CTkFont(size=12)).pack(side="left")
        self.start_num_var = ctk.StringVar(value="1")
        self.start_entry = ctk.CTkEntry(range_row, textvariable=self.start_num_var,
                                        width=60, height=32, font=ctk.CTkFont(size=12), justify="center")
        self.start_entry.pack(side="left", padx=(6, 16))

        ctk.CTkLabel(range_row, text="끝 번호", font=ctk.CTkFont(size=12)).pack(side="left")
        self.end_num_var = ctk.StringVar(value="")
        self.end_entry = ctk.CTkEntry(range_row, textvariable=self.end_num_var,
                                      width=60, height=32, font=ctk.CTkFont(size=12), justify="center")
        self.end_entry.pack(side="left", padx=(6, 0))

        self.total_label = ctk.CTkLabel(range_row, text="  (폴더를 선택해주세요)",
                                        font=ctk.CTkFont(size=11), text_color="gray")
        self.total_label.pack(side="left", padx=(8, 0))

        self._toggle_range()

        # 업로드 버튼
        self.start_btn = ctk.CTkButton(
            f4, text="▶   업로드 시작",
            font=ctk.CTkFont(size=15, weight="bold"),
            height=50, corner_radius=10,
            state="disabled",
            command=self._start_upload
        )
        self.start_btn.grid(row=9, column=0, columnspan=2, padx=16, pady=(0, 16), sticky="ew")

        # ── 진행 상황 ──
        fp = ctk.CTkFrame(sc)
        fp.grid(row=5, column=0, padx=24, pady=(12, 0), sticky="ew")
        fp.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(fp, text="진행 상황",
                     font=ctk.CTkFont(size=13, weight="bold")).grid(
            row=0, column=0, padx=16, pady=(14, 6), sticky="w")

        self.status_var = ctk.StringVar(value="대기 중...")
        self.status_label = ctk.CTkLabel(fp, textvariable=self.status_var,
                                         font=ctk.CTkFont(size=12), text_color="#555555", anchor="w")
        self.status_label.grid(row=1, column=0, padx=16, pady=(0, 6), sticky="ew")

        self.progress_bar = ctk.CTkProgressBar(fp, height=16, corner_radius=8)
        self.progress_bar.set(0)
        self.progress_bar.grid(row=2, column=0, padx=16, pady=(0, 6), sticky="ew")

        self.count_var = ctk.StringVar(value="")
        ctk.CTkLabel(fp, textvariable=self.count_var,
                     font=ctk.CTkFont(size=11), text_color="gray").grid(
            row=3, column=0, padx=16, pady=(0, 14), sticky="e")

        # ── 로그 ──
        fl = ctk.CTkFrame(sc)
        fl.grid(row=6, column=0, padx=24, pady=(12, 24), sticky="nsew")
        fl.grid_columnconfigure(0, weight=1)
        fl.grid_rowconfigure(1, weight=1)
        sc.grid_rowconfigure(6, weight=1)

        ctk.CTkLabel(fl, text="작업 로그",
                     font=ctk.CTkFont(size=13, weight="bold")).grid(
            row=0, column=0, padx=16, pady=(14, 6), sticky="w")

        self.log_text = ctk.CTkTextbox(fl, font=ctk.CTkFont(family="Courier New", size=11),
                                       state="disabled", wrap="word")
        self.log_text.grid(row=1, column=0, padx=16, pady=(0, 16), sticky="nsew")

    # ── 자동 업데이트 ──────────────────────────────
    def _show_update_done(self):
        """재시작 후 업데이트 완료 알림 표시."""
        if not os.path.exists(UPDATE_DONE_FILE):
            return
        try:
            with open(UPDATE_DONE_FILE, "r", encoding="utf-8") as f:
                info = json.load(f)
            os.remove(UPDATE_DONE_FILE)
            messagebox.showinfo(
                "업데이트 완료",
                f"업데이트가 완료되었습니다.\n\n"
                f"버전: {info['version']}\n"
                f"업데이트 시각: {info['updated_at']}"
            )
        except Exception:
            pass

    def _check_update(self):
        latest = fetch_latest_version()
        if latest and latest != VERSION:
            self.after(0, lambda: self._prompt_update(latest))

    def _prompt_update(self, latest: str):
        answer = messagebox.askyesno(
            "업데이트 알림",
            f"새 버전이 있습니다.\n\n현재: {VERSION}  →  최신: {latest}\n\n지금 업데이트할까요?"
        )
        if not answer:
            return

        tmp = tempfile.NamedTemporaryFile(suffix=".py", delete=False)
        tmp.close()

        self._set_status("업데이트 다운로드 중...", "#1565c0")
        ok = download_update(tmp.name)
        if ok:
            messagebox.showinfo("업데이트", "다운로드 완료!\n프로그램을 재시작합니다.")
            apply_update_and_restart(tmp.name, latest)
        else:
            os.unlink(tmp.name)
            messagebox.showerror("업데이트 실패", "다운로드 중 오류가 발생했습니다.\n수동으로 업데이트해주세요.")
            self._set_status("대기 중...", "#555555")

    # ── 헬퍼 ─────────────────────────────────────
    def _toggle_range(self):
        state = "disabled" if self.range_all_var.get() else "normal"
        self.start_entry.configure(state=state)
        self.end_entry.configure(state=state)

    def _browse_src(self):
        path = filedialog.askopenfilename(
            title="현금영수증 엑셀 파일 선택",
            filetypes=[("Excel 파일", "*.xlsx *.xls"), ("모든 파일", "*.*")]
        )
        if path:
            self.file_var.set(path)
            default_out = os.path.join(os.path.dirname(path), "현금영수증_분할")
            self.out_dir_var.set(default_out)
            self.split_files = []
            self.split_result_var.set("")
            self.total_label.configure(text="  (폴더를 선택해주세요)", text_color="gray")
            self.start_btn.configure(state="disabled")
            self.split_info_frame.grid_remove()

    def _load_existing_folder(self):
        """분할 파일이 있는 폴더를 선택해서 업로드 목록으로 불러옴"""
        folder = filedialog.askdirectory(title="업로드할 분할 파일 폴더 선택")
        if not folder:
            return
        self._apply_upload_folder(folder)

    def _apply_upload_folder(self, folder: str):
        """폴더를 스캔해서 분할 파일 목록을 Step 4에 반영"""
        import glob
        files = sorted(
            glob.glob(os.path.join(folder, "현금영수증_분할_*.xlsx")),
            key=lambda f: int(os.path.splitext(os.path.basename(f))[0].split("_")[-1])
        )
        if not files:
            messagebox.showwarning("알림", "선택한 폴더에 분할 파일이 없습니다.\n(현금영수증_분할_N.xlsx 형식의 파일을 찾을 수 없음)")
            return

        n = len(files)
        self.split_files = files
        self.split_save_dir = folder
        self.upload_dir_var.set(folder)
        self.total_label.configure(text=f"  (전체 {n}개)", text_color="#2e7d32")
        self.start_num_var.set("1")
        self.end_num_var.set(str(n))
        self.start_btn.configure(state="normal")
        self.step3_open_btn.configure(state="normal")

    def _open_split_folder(self):
        folder = getattr(self, "split_save_dir", None)
        if folder and os.path.exists(folder):
            import subprocess, sys
            if sys.platform == "win32":
                os.startfile(folder)
            elif sys.platform == "darwin":
                subprocess.run(["open", folder])
            else:
                subprocess.run(["xdg-open", folder])

    def _browse_out(self):
        path = filedialog.askdirectory(title="분할 파일 저장 폴더 선택")
        if path:
            self.out_dir_var.set(path)

    def _log(self, text, color=None):
        self.msg_queue.put(("log", text, color))

    def _set_status(self, text, color="#555555"):
        self.msg_queue.put(("status", text, color))

    def _set_progress(self, cur, total):
        self.msg_queue.put(("progress", cur, total))

    def _poll_queue(self):
        try:
            while True:
                msg = self.msg_queue.get_nowait()
                kind = msg[0]
                if kind == "log":
                    _, text, _ = msg
                    ts = datetime.now().strftime("%H:%M:%S")
                    self.log_text.configure(state="normal")
                    self.log_text.insert("end", f"[{ts}] {text}\n")
                    self.log_text.see("end")
                    self.log_text.configure(state="disabled")
                elif kind == "status":
                    _, text, color = msg
                    self.status_var.set(text)
                    self.status_label.configure(text_color=color)
                elif kind == "progress":
                    _, cur, total = msg
                    self.progress_bar.set(cur / total if total > 0 else 0)
                    self.count_var.set(f"{cur} / {total} 파일 완료")
                elif kind == "split_done":
                    _, files = msg
                    n = len(files)
                    if n > 0:
                        self.split_result_var.set(f"✅  총 {n}개 파일 생성됨")
                        folder = os.path.dirname(files[0])
                        self.info_folder_var.set(f"📁  {folder}")
                        self.split_info_frame.grid()
                        self._apply_upload_folder(folder)
                    self.split_btn.configure(state="normal", text="파일 분할")
                    self.running = False
                elif kind == "done":
                    self.start_btn.configure(state="normal" if self.split_files else "disabled",
                                             text="▶   업로드 시작")
                    self.running = False
        except queue.Empty:
            pass
        self.after(100, self._poll_queue)

    # ── 분할 ─────────────────────────────────────
    def _start_split(self):
        if self.running:
            return
        file_path = self.file_var.get().strip()
        out_dir = self.out_dir_var.get().strip()
        if not file_path:
            messagebox.showwarning("알림", "원본 엑셀 파일을 선택해주세요.")
            return
        if not os.path.isfile(file_path):
            messagebox.showwarning("알림", "선택한 파일이 존재하지 않습니다.")
            return
        if not out_dir:
            messagebox.showwarning("알림", "분할 파일 저장 폴더를 지정해주세요.")
            return

        self.running = True
        self.split_btn.configure(state="disabled", text="분할 중...")
        self.split_result_var.set("")
        self.split_files = []
        self.start_btn.configure(state="disabled")

        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")
        self.progress_bar.set(0)
        self.count_var.set("")

        threading.Thread(target=self._run_split, args=(file_path, out_dir), daemon=True).start()

    def _run_split(self, file_path, out_dir):
        self._set_status("엑셀 파일 분할 중...", "#1565c0")
        self._log("[ 파일 분할 ] 시작")
        try:
            files = split_excel(file_path, out_dir)
            self._log(f"✅ 분할 완료 → {len(files)}개 파일")
            self._log(f"  저장 위치: {out_dir}")
            self._set_status(f"✅ 분할 완료  |  총 {len(files)}개 파일 → {out_dir}", "#2e7d32")
            self.msg_queue.put(("split_done", files))
        except Exception as e:
            self._log(f"❌ 분할 실패: {e}")
            self._set_status("❌ 분할 실패", "#c62828")
            self.msg_queue.put(("split_done", []))

    # ── 업로드 ───────────────────────────────────
    def _start_upload(self):
        if self.running:
            return
        if not self.split_files:
            messagebox.showwarning("알림", "먼저 파일 분할을 실행해주세요.")
            return

        user_id = self.id_var.get().strip()
        user_pw = self.pw_var.get().strip()
        if not user_id or not user_pw:
            messagebox.showwarning("알림", "아이디와 비밀번호를 입력해주세요.")
            return

        total = len(self.split_files)
        if not self.range_all_var.get():
            try:
                start_n = int(self.start_num_var.get())
                end_n = int(self.end_num_var.get())
                if start_n < 1 or end_n < start_n or start_n > total:
                    raise ValueError
            except ValueError:
                messagebox.showwarning("알림", f"업로드 범위를 올바르게 입력해주세요.\n(1 ~ {total} 사이, 시작 ≤ 끝)")
                return

        save_config(user_id, user_pw)

        if self.range_all_var.get():
            upload_files = self.split_files
        else:
            s = int(self.start_num_var.get())
            e = min(int(self.end_num_var.get()), total)
            upload_files = self.split_files[s - 1:e]

        self.running = True
        self.start_btn.configure(state="disabled", text="⏳  업로드 중...")

        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")
        self.progress_bar.set(0)
        self.count_var.set("")

        threading.Thread(target=self._run_upload, args=(upload_files, user_id, user_pw), daemon=True).start()

    def _run_upload(self, upload_files, user_id, user_pw):
        driver = None
        upload_count = len(upload_files)
        try:
            self._set_status("크롬 드라이버 준비 중... (처음 실행 시 다운로드)", "#1565c0")
            self._log("[ 업로드 ] 크롬 드라이버 초기화 중...")
            self._log("  (처음 실행 시 드라이버 다운로드로 시간이 걸릴 수 있습니다)")
            try:
                opts = Options()
                service = Service(ChromeDriverManager().install())
                driver = webdriver.Chrome(service=service, options=opts)
                driver.minimize_window()
            except Exception as e:
                self._log(f"❌ 크롬 드라이버 초기화 실패: {e}")
                self._log("  → 크롬 브라우저가 설치되어 있는지 확인해주세요")
                self._set_status("❌ 실패: 크롬 드라이버 오류", "#c62828")
                self.msg_queue.put(("done",))
                return
            self._log("✅ 크롬 드라이버 초기화 완료")

            self._set_status("토스페이먼츠 로그인 중...", "#1565c0")
            self._log("  로그인 시도 중...")
            try:
                driver.get(LOGIN_URL)
                WebDriverWait(driver, 15).until(
                    EC.presence_of_element_located((By.ID, "user_id"))
                ).send_keys(user_id)
                driver.find_element(By.ID, "user_pw").send_keys(user_pw)
                driver.find_element(By.ID, "login").click()
                WebDriverWait(driver, 15).until(EC.url_contains("main"))
                self._log("✅ 로그인 완료")
            except Exception as e:
                self._log(f"❌ 로그인 실패: {e}")
                self._log("  → 아이디/비밀번호를 다시 확인해주세요")
                self._set_status("❌ 실패: 로그인 오류", "#c62828")
                self.msg_queue.put(("done",))
                return

            self._set_status(f"파일 업로드 중... (0 / {upload_count})", "#1565c0")
            self._log(f"[ 업로드 시작 ] {upload_count}개 파일")
            for idx, file in enumerate(upload_files, 1):
                fname = os.path.basename(file)
                self._set_status(f"업로드 중: {fname}  ({idx} / {upload_count})", "#1565c0")
                self._log(f"  [{idx}/{upload_count}] {fname} 업로드 중...")
                try:
                    driver.get(UPLOAD_URL)
                    WebDriverWait(driver, 15).until(
                        EC.presence_of_element_located((By.ID, "uploadFile"))
                    )
                    driver.find_element(By.ID, "uploadFile").send_keys(os.path.abspath(file))
                    time.sleep(4)
                    self._log(f"  ✅ [{idx}/{upload_count}] {fname} 업로드 완료")
                    self._set_progress(idx, upload_count)
                except Exception as e:
                    self._log(f"  ❌ [{idx}/{upload_count}] {fname} 업로드 실패: {e}")
                    self._log("  → 오류 발생으로 나머지 업로드를 중단합니다")
                    self._set_status(f"❌ {idx}번째 파일 오류로 중단됨 → {fname}", "#c62828")
                    self._set_progress(idx, upload_count)
                    self.msg_queue.put(("done",))
                    return

            self._set_status(f"✅ 완료!  {upload_count}개 파일 업로드 성공", "#2e7d32")
            self._log("")
            self._log(f"🎉 완료!  {upload_count}개 파일 업로드 성공")
            self.after(0, lambda n=upload_count: messagebox.showinfo(
                "완료", f"{n}개 파일을 모두 업로드했습니다!\n\n창을 닫아도 됩니다."
            ))

        except Exception as e:
            self._log(f"❌ 예기치 못한 오류: {e}")
            self._set_status("❌ 예기치 못한 오류 발생", "#c62828")
        finally:
            if driver:
                try:
                    driver.quit()
                except Exception:
                    pass
            self.msg_queue.put(("done",))


if __name__ == "__main__":
    app = App()
    app.mainloop()
